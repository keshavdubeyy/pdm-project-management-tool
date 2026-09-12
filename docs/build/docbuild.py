"""Shared helpers for building the PDM Word deliverables."""

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

# ---------------------------------------------------------------- palette
NAVY = "1F3864"
NAVY_RGB = RGBColor(0x1F, 0x38, 0x64)
STEEL = "2E5C8A"
STEEL_RGB = RGBColor(0x2E, 0x5C, 0x8A)
INK_RGB = RGBColor(0x1A, 0x1A, 0x1A)
GREY_RGB = RGBColor(0x55, 0x5F, 0x6D)

FILL_HEADER = "1F3864"
FILL_SUBHEAD = "D9E2F0"
FILL_ZEBRA = "F4F6F9"
FILL_WHITE = "FFFFFF"
FILL_RED = "FBD9D9"
FILL_AMBER = "FCE9CC"
FILL_GREEN = "D8EEDB"
FILL_BLUE = "DCE8F6"
FILL_GREY = "ECECEC"
FILL_PURPLE = "E6DCF0"

TEXT_RED = RGBColor(0x9B, 0x1C, 0x1C)
TEXT_AMBER = RGBColor(0x7A, 0x45, 0x00)
TEXT_GREEN = RGBColor(0x1B, 0x5E, 0x20)
TEXT_NAVY = NAVY_RGB

PAGE_W = Cm(21.0)
CONTENT_W = Cm(17.0)   # A4 width minus 2cm margins each side
MAX_FIG_H = Cm(19.0)


# ---------------------------------------------------------------- low level
def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:shd")):
        tcPr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _cell_margins(table, top=60, bottom=60, left=110, right=110):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def _borders(table, colour="BFC9D9", size=4):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), colour)
        borders.append(el)
    tblPr.append(borders)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _no_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    trPr.append(el)


def _set_col_widths(table, widths_cm):
    """Fixed layout. Word and LibreOffice both honour tblGrid, so write it."""
    table.autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    for old in tblPr.findall(qn("w:tblLayout")):
        tblPr.remove(old)
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblPr.append(layout)

    twips = [int(round(w * 567)) for w in widths_cm]
    for old in tbl.findall(qn("w:tblGrid")):
        tbl.remove(old)
    grid = OxmlElement("w:tblGrid")
    for t in twips:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(t))
        grid.append(gc)
    tbl.insert(list(tbl).index(tblPr) + 1, grid)

    for row in table.rows:
        for idx, t in enumerate(twips):
            if idx < len(row.cells):
                tcPr = row.cells[idx]._tc.get_or_add_tcPr()
                for old in tcPr.findall(qn("w:tcW")):
                    tcPr.remove(old)
                tcW = OxmlElement("w:tcW")
                tcW.set(qn("w:w"), str(t))
                tcW.set(qn("w:type"), "dxa")
                tcPr.append(tcW)


# ---------------------------------------------------------------- document
def new_document(font="Calibri"):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(1.9)
    sec.left_margin = Cm(2.0)
    sec.right_margin = Cm(2.0)
    sec.header_distance = Cm(1.1)
    sec.footer_distance = Cm(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK_RGB
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.18
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), font)
    rfonts.set(qn("w:cs"), font)

    for name, size, colour, before, after, bold in (
        ("Heading 1", 16, NAVY_RGB, 20, 8, True),
        ("Heading 2", 12.5, STEEL_RGB, 14, 5, True),
        ("Heading 3", 11, NAVY_RGB, 10, 4, True),
    ):
        st = doc.styles[name]
        st.font.name = font
        st.font.size = Pt(size)
        st.font.color.rgb = colour
        st.font.bold = bold
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), font)
    return doc


def footer_page_numbers(doc, left_text):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.text = ""
    p.paragraph_format.tab_stops.clear_all()
    from docx.enum.text import WD_TAB_ALIGNMENT
    p.paragraph_format.tab_stops.add_tab_stop(Cm(17.0), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(left_text + "\t")
    r.font.size = Pt(8)
    r.font.color.rgb = GREY_RGB
    _add_field(p, "PAGE", size=8, colour=GREY_RGB)


def _add_field(paragraph, instr, size=9, colour=GREY_RGB):
    run = paragraph.add_run()
    run.font.size = Pt(size)
    run.font.color.rgb = colour
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = f" {instr} "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr_el)
    run._r.append(fld_sep)
    run._r.append(fld_end)


def contents(doc, entries, page_map=None):
    """entries = [(level, text), ...]. Rendered as a plain list with dotted
    leaders and page numbers once page_map is known."""
    from docx.enum.text import WD_TAB_ALIGNMENT
    for level, text in entries:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4 if level == 1 else 2)
        p.paragraph_format.tab_stops.clear_all()
        p.paragraph_format.tab_stops.add_tab_stop(
            Cm(17.0), WD_TAB_ALIGNMENT.RIGHT, 1)  # 1 = dotted leader
        if level == 2:
            p.paragraph_format.left_indent = Cm(0.8)
        r = p.add_run(text)
        r.font.size = Pt(11 if level == 1 else 10)
        r.font.bold = level == 1
        r.font.color.rgb = NAVY_RGB if level == 1 else INK_RGB
        num = str(page_map.get(text, "")) if page_map else ""
        r2 = p.add_run("\t" + num)
        r2.font.size = Pt(10)
        r2.font.color.rgb = GREY_RGB if level == 2 else NAVY_RGB
        r2.font.bold = level == 1


def page_map_from_docx(docx_path, entries, workdir, skip_pages=2):
    """Render to PDF once and read back which page each heading landed on."""
    import glob, os, re, shutil, subprocess
    tmp = os.path.join(workdir, "_pagination")
    shutil.rmtree(tmp, ignore_errors=True)
    os.makedirs(tmp, exist_ok=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf",
                    "--outdir", tmp, docx_path],
                   check=True, capture_output=True)
    pdf = glob.glob(os.path.join(tmp, "*.pdf"))[0]
    txt = subprocess.run(["pdftotext", "-layout", pdf, "-"],
                         check=True, capture_output=True, text=True).stdout
    pages = txt.split("\f")
    mapping = {}
    for _, text in entries:
        needle = re.sub(r"\s+", " ", text).strip().lower()
        for idx, page in enumerate(pages, start=1):
            if idx <= skip_pages:      # cover and contents repeat every heading
                continue
            flat = re.sub(r"\s+", " ", page).lower()
            if needle in flat:
                mapping[text] = idx
                break
    return mapping


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------- blocks
def cover(doc, title, subtitle, meta_rows, strapline=None):
    band = doc.add_paragraph()
    band.paragraph_format.space_before = Pt(60)
    band.paragraph_format.space_after = Pt(2)
    r = band.add_run(strapline or "")
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = STEEL_RGB

    t = doc.add_paragraph()
    t.paragraph_format.space_after = Pt(4)
    r = t.add_run(title)
    r.font.size = Pt(27)
    r.font.bold = True
    r.font.color.rgb = NAVY_RGB

    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(10)
    pPr = rule._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), NAVY)
    bdr.append(bottom)
    pPr.append(bdr)

    s = doc.add_paragraph()
    s.paragraph_format.space_after = Pt(26)
    r = s.add_run(subtitle)
    r.font.size = Pt(12.5)
    r.font.color.rgb = GREY_RGB

    tbl = doc.add_table(rows=0, cols=2)
    _borders(tbl, colour="FFFFFF", size=2)
    _cell_margins(tbl, top=40, bottom=40, left=0, right=80)
    for k, v in meta_rows:
        row = tbl.add_row()
        c0, c1 = row.cells
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(10)
        r0.font.color.rgb = NAVY_RGB
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.size = Pt(10)
    _set_col_widths(tbl, [5.0, 12.0])


def para(doc, text, size=10.5, bold=False, italic=False, colour=None, space_after=7,
         align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    if colour is not None:
        r.font.color.rgb = colour
    return p


def rich_para(doc, parts, size=10.5, space_after=7):
    """parts = [(text, bold, colour_or_None), ...]"""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    for text, bold, colour in parts:
        r = p.add_run(text)
        r.font.size = Pt(size)
        r.font.bold = bold
        if colour is not None:
            r.font.color.rgb = colour
    return p


def bullets(doc, items, size=10.5):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        if isinstance(it, tuple):
            lead, rest = it
            r = p.add_run(lead)
            r.font.bold = True
            r.font.size = Pt(size)
            r2 = p.add_run(rest)
            r2.font.size = Pt(size)
        else:
            r = p.add_run(it)
            r.font.size = Pt(size)


def numbered(doc, items, size=10.5):
    for it in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = Cm(-0.45)
        r = p.add_run(it)
        r.font.size = Pt(size)


def callout(doc, text, fill=FILL_BLUE, text_colour=TEXT_NAVY, label=None):
    tbl = doc.add_table(rows=1, cols=1)
    _borders(tbl, colour="FFFFFF", size=2)
    _cell_margins(tbl, top=110, bottom=110, left=160, right=160)
    cell = tbl.rows[0].cells[0]
    _shade(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if label:
        r = p.add_run(label + "  ")
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = text_colour
    r = p.add_run(text)
    r.font.size = Pt(10)
    r.font.color.rgb = text_colour
    _set_col_widths(tbl, [17.0])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def table(doc, headers, rows, widths, fills=None, header_fill=FILL_HEADER,
          font_size=9, zebra=True, align_center_cols=()):
    """rows: list of list[str]. fills: dict {(row_idx, col_idx): hex} or
    {row_idx: hex} for a whole row."""
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    _borders(tbl)
    _cell_margins(tbl)
    hdr = tbl.rows[0]
    _repeat_header(hdr)
    for i, h in enumerate(headers):
        cell = hdr.cells[i]
        _shade(cell, header_fill)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        if i in align_center_cols:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(font_size)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for ri, row_vals in enumerate(rows):
        row = tbl.add_row()
        _no_split(row)
        base = FILL_ZEBRA if (zebra and ri % 2 == 1) else FILL_WHITE
        for ci, val in enumerate(row_vals):
            cell = row.cells[ci]
            fill = base
            colour = None
            bold = False
            if fills:
                if ri in fills and not isinstance(fills[ri], dict):
                    fill = fills[ri]
                key = (ri, ci)
                if key in fills:
                    spec = fills[key]
                    if isinstance(spec, tuple):
                        fill, colour, bold = spec
                    else:
                        fill = spec
            _shade(cell, fill)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            if ci in align_center_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            lines = str(val).split("\n")
            for li, line in enumerate(lines):
                target = p if li == 0 else cell.add_paragraph()
                if li:
                    target.paragraph_format.space_before = Pt(1)
                    target.paragraph_format.space_after = Pt(0)
                    target.paragraph_format.line_spacing = 1.08
                    if ci in align_center_cols:
                        target.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = target.add_run(line)
                r.font.size = Pt(font_size)
                r.font.bold = bold or (li == 0 and len(lines) > 1)
                if colour is not None:
                    r.font.color.rgb = colour
    _set_col_widths(tbl, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return tbl


def figure(doc, path, caption, max_width_cm=17.0, max_height_cm=18.5):
    w_px, h_px = Image.open(path).size
    width = max_width_cm
    height = width * h_px / w_px
    if height > max_height_cm:
        height = max_height_cm
        width = height * w_px / h_px
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(path, width=Cm(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(12)
    r = cap.add_run(caption)
    r.font.size = Pt(8.5)
    r.font.italic = True
    r.font.color.rgb = GREY_RGB


def legend(doc, entries):
    """entries = [(fill_hex, label), ...] rendered as one coloured strip."""
    tbl = doc.add_table(rows=1, cols=len(entries))
    _borders(tbl, colour="FFFFFF", size=2)
    _cell_margins(tbl, top=50, bottom=50, left=80, right=80)
    for i, (fill, label) in enumerate(entries):
        cell = tbl.rows[0].cells[i]
        _shade(cell, fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label)
        r.font.size = Pt(8)
        r.font.bold = True
    w = 17.0 / len(entries)
    _set_col_widths(tbl, [w] * len(entries))
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
